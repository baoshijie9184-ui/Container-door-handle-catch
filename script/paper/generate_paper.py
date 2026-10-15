# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('docs/集装箱门把手智能检测与多智能体协同系统设计与实现完整版.docx')
doc = Document()
s = doc.sections[0]
s.top_margin = Inches(.8); s.bottom_margin = Inches(.75); s.left_margin = Inches(.9); s.right_margin = Inches(.9)
for name, size in [('Normal',10.5),('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
    st=doc.styles[name]; st.font.name='SimSun' if name == 'Normal' else 'SimHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体' if name == 'Normal' else '黑体'); st._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].paragraph_format.line_spacing=1.35; doc.styles['Normal'].paragraph_format.space_after=Pt(6)

def para(t, style=None, center=False):
    x=doc.add_paragraph(style=style); x.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT; x.add_run(t); return x
def heading(t, level=1): doc.add_heading(t, level=level)
def page(): doc.add_page_break()
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def tbl(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'1F4E78'); c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    doc.add_paragraph()

para('集装箱门把手智能检测与多智能体协同系统设计与实现','Title',True)
para('基于 Spring Boot 模块化单体与 AgentScope Java 的企业级应用研究',center=True)
for _ in range(6): para('')
for x in ['论文类型：系统设计与工程实现','研究对象：集装箱门把手工业缺陷检测平台','技术路线：Java 17、Spring Boot、MySQL、OSS、Kafka、AgentScope Java 2.0.1','版本日期：2026 年 9 月']: para(x,center=True)
page()

heading('摘要')
para('随着工业视觉检测任务规模不断扩大，检测人员需要在图片上传、模型推理、结果复核、质量审批和历史追溯之间频繁切换。传统页面往往将上传、推理和人工审核耦合在一起，难以处理多人协同、重复图片、结果版本和智能辅助等问题。本文以集装箱门把手缺陷检测平台为对象，提出一种面向单企业场景的模块化单体与多智能体协同架构。系统以检测任务为核心业务聚合，以对象存储保存图片，以消息事件驱动远程推理，以主管审批控制公共数据发布，并使用乐观锁、状态机和 Outbox 事件保证多用户并发一致性。')
para('在智能化方面，系统引入 AgentScope Java 2.0.1，设计员工助手、检测分析、缺陷解释、报告草稿、知识检索、主管质量分析和重复工单等专用智能体。业务数据按公司、部门、用户和资源范围隔离，Agent 对话则按用户和会话工作区独立存储，做到共享检测事实而不共享私人对话。针对本地离线上传导致的重复图片问题，系统通过 SHA-256 指纹、重复工单、左右双面板对比和主管确认建立可审计的解决链路。')
para('关键词：工业视觉检测；模块化单体；AgentScope Java；多智能体；乐观锁；重复图片；质量审批')
page()

heading('Abstract')
para('This paper presents a modular-monolith and multi-agent architecture for a container door-handle defect inspection platform. Detection tasks are treated as business aggregates, images are stored in object storage, remote inference is driven by events, and public records are released only after supervisor approval. Optimistic locking, state machines, and an outbox pattern provide consistency for concurrent users. AgentScope Java 2.0.1 organizes employee and supervisor assistants, inspection analysis, defect explanation, report drafting, knowledge retrieval, and duplicate-ticket comparison. Business facts are shared according to department scope, while agent conversations and workspaces are isolated by user and session.')
para('Keywords: industrial inspection; modular monolith; AgentScope Java; multi-agent system; optimistic locking; duplicate image; quality approval')
page()

heading('目录')
for x in ['摘要','Abstract','第1章 绪论','第2章 需求分析与业务规则','第3章 系统总体架构设计','第4章 检测业务流程与权限模型','第5章 多用户共享数据与一致性设计','第6章 AgentScope 多智能体设计','第7章 数据库与接口设计','第8章 前端交互与复查对比设计','第9章 部署配置与可观测性','第10章 测试方案与质量保证','第11章 系统演进与微服务拆分','第12章 结论与展望','参考文献']: para(x)
page()

chapters = [
('第1章 绪论',[('1.1 研究背景','集装箱门把手是运输装备中的关键结构件，其表面缺陷、锈蚀、裂纹和装配异常会影响安全与使用寿命。工业企业通常以文件夹或对象存储的形式积累大量图片，再通过视觉模型进行批量推理。随着数据量增加，系统必须同时解决上传可靠性、异步推理、多人任务协作、重复检测、人工复核和质量追溯等问题。'),('1.2 研究内容','本文构建以业务模块为边界的 Java 模块化单体，设计检测任务从创建、处理、提交、审批到公共发布的状态机，使用图片指纹和工作单元管理最多 5000 张图片的批量检测，并构建员工与主管两套多智能体体系。'),('1.3 论文组织','全文依次讨论需求、架构、流程、并发一致性、智能体、数据库接口、前端、部署测试和演进路线。')]),
('第2章 需求分析与业务规则',[('2.1 角色需求','普通检测员工负责上传图片、执行本人任务、提交审批；检测主管负责派发工作、复查、处理重复图片并发布公共检测记录；系统管理员维护账号、角色和基础设施配置。'),('2.2 检测业务规则','检测任务采用私有处理、主管审批、部门公共只读的发布模型。员工创建任务后，任务及图片只对创建者和授权主管可见。主管通过后生成公共检测记录，检测部门员工均可查询，但公共记录修改权只授予主管。'),('2.3 图片批次与重复规则','上传文件夹后先预检查，按 SHA-256 判断历史状态。员工可选择仅检测未成功图片或全部重新检测。离线批次在主管池按指纹聚合，重复内容生成重复工单。')]),
('第3章 系统总体架构设计',[('3.1 模块化单体','系统包含 dependencies、common、framework、contracts、业务模块和 server。Server 只负责启动、配置和依赖装配；检测、系统、基础设施、模型、设备和 AI 能力分别归属业务模块。'),('3.2 逻辑分层','模块内部采用 Controller、Service、Mapper、Entity、DTO/VO 和 API 分层。Controller 不访问数据库，跨模块调用只传递 DTO 或事件。'),('3.3 技术链路','前端通过 REST 和 AG-UI 交互；图片通过 OSS 预签名 URL 上传；检测任务发布 Kafka 事件；远程检测服务回写结果；主管审批后发布公共资源。')]),
('第4章 检测业务流程与权限模型',[('4.1 任务状态机','任务状态包括 DRAFT、UPLOADING、RUNNING、READY_FOR_REVIEW、RETURNED、APPROVED 和 PUBLIC。状态转换由 Service 控制，失败任务可重试。'),('4.2 员工任务与公共资源','员工以 created_by 作为默认所有者，员工 A 不能操作员工 B 的任务。审批通过后公共查询按部门返回记录，员工只能查询和查看。'),('4.3 权限校验','普通员工拥有任务查询、结果查询、复核提交和草稿写入权限；主管增加审批、复检和报告发布权限。身份从 Spring Security 上下文取得。')]),
('第5章 多用户共享数据与一致性设计',[('5.1 共享与私有','共享业务数据包括审批通过的检测记录、图片引用和统计；私有数据包括 Agent 对话、未提交草稿和工作区临时文件。'),('5.2 乐观锁','草稿、公共记录和重复工单包含 version。更新携带 expectedVersion，版本不一致返回 409，前端重新加载并提示合并。'),('5.3 事件同步','Service 事务同时写业务表和 event_outbox，发布器按 eventId 幂等发送事件，客户端收到资源变更后重新查询最新版本。')]),
('第6章 AgentScope 多智能体设计',[('6.1 员工体系','employee-assistant 编排 inspection-analysis、defect-vision、report-draft、knowledge 和 safety-router。编排器只负责意图识别和结果汇总。'),('6.2 主管体系','supervisor-assistant 编排 quality-overview、review-approval、reinspection-planner、report-publisher 和 compliance。审批和发布仍需人工确认。'),('6.3 工具与 Skills','工具使用稳定命名和 DTO 参数，身份由 RuntimeContext 注入。Skills 记录角色、触发条件、工具白名单、输出 Schema 和版本。')]),
('第7章 数据库与接口设计',[('7.1 核心表','detection_task 保存任务聚合；detection_image_fingerprint 保存 SHA-256 和状态；detection_work_unit 管理最多 5000 张图片；detection_work_unit_image 追踪图片；detection_duplicate_ticket 保存重复工单。'),('7.2 链路字段','公共记录保存图片 ID、原文件名、对象 Key、SHA-256、员工 ID、姓名快照、模型版本、检测时间、审批人和结果版本。'),('7.3 接口设计','员工使用 batch-precheck、work-units/apply、work-units/mine 和 submit；主管使用 pending、dispatch、recall、duplicate-tickets 和 resolve。')]),
('第8章 前端交互与复查对比设计',[('8.1 员工工作台','工作台包含模型选择、文件夹上传、批量预检查、策略选择、进度和提交审批。'),('8.2 主管复查页面','重复工单采用左右双面板，分别显示两个来源员工、文件夹、检测时间和图片。两侧共享 currentIndex，上一张和下一张同步切换。'),('8.3 对比结论','主管可采用左侧、右侧或要求重新检测，提交时携带 expectedVersion 和对比说明。')]),
('第9章 部署配置与可观测性',[('9.1 环境配置','仓库提交 .env.example，真实 .env 由本地或部署平台维护。AgentScope 最小配置为 AGENTSCOPE_AI_ENABLED、DASHSCOPE_API_KEY 和 AGENTSCOPE_AI_MODEL。'),('9.2 部署结构','script 按 docker、jenkins、shell 管理脚本；MySQL 迁移归档在 sql/mysql；Server 使用非 root 用户和 Actuator 健康检查。'),('9.3 观测指标','记录 traceId、userId、sessionId、runId、agentId、skillVersions、toolName、resourceId、latency 和 policyDecision。')]),
('第10章 测试方案与质量保证',[('10.1 测试层次','单元测试覆盖状态机和指纹；Web 测试覆盖权限和错误码；集成测试覆盖 MySQL、Redis、OSS 和 Kafka；端到端测试覆盖上传到公共发布。'),('10.2 并发测试','验证两个用户编辑同一草稿时一个成功、一个返回 409；两个主管处理同一工单时只有一个版本更新成功。'),('10.3 安全测试','验证修改 sessionId、taskId、userId 不能越权，工作区拒绝绝对路径和 ..，日志不泄露密钥。')]),
('第11章 系统演进与微服务拆分',[('11.1 模块边界','检测模块拥有检测数据，AI 模块拥有 Agent、工具和会话，contracts 保存 DTO 和事件契约。'),('11.2 拆分候选','AI、图片对象服务、GPU 推理服务和质量审批服务可在负载与团队边界稳定后拆分。'),('11.3 拆分前检查','消除 Entity 泄漏，稳定事件 Schema 和错误码，清理隐式跨模块事务，实施灰度和可回滚部署。')]),
('第12章 结论与展望',[('12.1 结论','本文提出面向单企业、多用户协同的模块化单体和 AgentScope 多智能体设计。任务私有处理、主管审批发布和公共记录只读解决了共享与责任边界问题。'),('12.2 展望','后续完善图片级推理状态、长期记忆、模型路由、成本控制和服务拆分，并持续评估 Agent 输出事实一致性与人工审批效率。')])]

detail_map = {
'第1章 绪论': ['从工程实践看，检测系统的难点并不只在模型精度，还在于把模型输出纳入可解释、可审批、可追溯的业务流程。若图片来源、操作者和模型版本缺少关联，即使算法结果准确，也无法支撑质量责任认定。', '本研究采用案例驱动的方法，将现有代码结构、检测接口和实际操作流程作为设计依据，再把重复图片和多人并发作为重点约束进行抽象。研究成果既服务当前项目，也为后续服务拆分提供边界。'],
'第2章 需求分析与业务规则': ['需求分析将系统参与者、资源状态和操作权限分开建模。员工关注任务完成效率，主管关注质量判断和责任追溯，系统则必须保证两者不会因共享数据而相互覆盖。', '本项目将审批通过作为公共可见性的唯一闸门，避免未复核的模型结果直接影响其他员工。对于本地离线上传无法实时去重的情况，系统采用事后指纹聚合和重复工单补偿，保证业务可继续而不牺牲审计完整性。'],
'第3章 系统总体架构设计': ['模块化单体在当前阶段能够降低部署和运维复杂度，同时通过 Maven 模块和 API Facade 保持领域边界。Server 作为空壳容器只装配模块，未来可以将 detection 或 ai 模块替换为远程客户端，而不改变前端业务语义。', '事件驱动用于隔离同步请求和长耗时推理。用户提交任务后立即获得任务编号，推理服务异步消费事件并回写进度，页面通过查询或推送获得状态。这种设计提高了大批量图片处理的韧性。'],
'第4章 检测业务流程与权限模型': ['权限模型采用功能权限与数据范围的交集。即使员工拥有结果查询权限，也只能查看自己的未审批任务和部门内已发布记录；主管拥有审批权限，也必须受所属部门范围约束。', '状态机把业务规则固化在服务层，而不是散落在前端按钮中。任何接口、消息消费者或 Agent 工具都必须调用同一组状态转换方法，避免绕过页面形成非法状态。'],
'第5章 多用户共享数据与一致性设计': ['乐观锁适用于员工和主管大多数编辑场景，因为页面编辑时间较长而真正提交时间很短。版本冲突后让用户看到差异，比长时间持有数据库锁更符合企业系统的可用性要求。', '对于检测结果，系统进一步采用不可变事实和追加式复核。这样 A 和 B 的判断都可以保留，主管的最终选择成为新的业务事实，历史记录不因后续修订而消失。'],
'第6章 AgentScope 多智能体设计': ['多智能体拆分依据是业务职责而不是模型数量。一个 Agent 负责一个可解释的任务边界，编排器负责选择和汇总，工具负责执行和授权。这样可以分别评估每个 Agent 的准确性和安全性。', 'AgentScope 2.0.1 的 HarnessAgent 作为运行载体，AG-UI 负责流式交互。模型输出不能直接写库，所有写入必须通过工具、权限策略、确认令牌和事务服务完成。'],
'第7章 数据库与接口设计': ['数据表设计围绕检测任务、图片指纹、工作单元和重复工单展开。任务记录业务聚合，指纹表解决重复识别，工作单元解决人员分配，重复工单解决冲突结果的人工裁决。', '接口设计遵循请求模型和响应模型分离原则。图片地址使用短期签名 URL，敏感字段不出现在日志中；写接口统一支持幂等键和版本号，为重试和并发提供基础。'],
'第8章 前端交互与复查对比设计': ['员工页面首先呈现预检查统计，让用户在启动模型前知道哪些图片会被跳过、重试或重新检测。主管页面以重复工单为中心，左右面板使用同一索引，从视觉上减少人工在两个窗口之间切换造成的误判。', '对比结论不是简单的“选左边或右边”，而是带有说明、操作者、版本和后续动作。选择重新检测时，系统生成新的运行编号；选择一侧时，系统保留另一侧作为历史来源。'],
'第9章 部署配置与可观测性': ['环境变量采用示例文件与真实文件分离的方式，仓库只保存占位符。数据库密码、OSS 密钥、DashScope Key 和 Langfuse Secret 由本地或部署平台注入，应用配置只读取变量。', '可观测性同时覆盖传统检测任务和 Agent 运行。通过统一 traceId 可以关联一次上传、一次推理、一次审批和一次工具调用，为定位慢请求、错误结果和权限拒绝提供依据。'],
'第10章 测试方案与质量保证': ['检测系统的测试重点是状态和权限，而不是只测试 Controller 返回 200。测试需要构造重复提交、断线重试、版本冲突、任务回收和跨员工访问等场景，验证系统在异常路径下仍然保持数据边界。', 'Agent 测试采用固定输入集和工具调用断言，检查模型是否选择正确工具、是否引用真实事实、是否在需要确认时停止。生产环境还应持续收集人工修订结果作为评估样本。'],
'第11章 系统演进与微服务拆分': ['拆分微服务前必须先稳定模块内的所有权和契约。若检测模块仍然依赖 AI 模块的实体或共享事务，拆分后会产生隐式网络调用和难以恢复的一致性问题。', '本项目优先把高资源消耗的推理、图片对象管理和 Agent 运行作为候选拆分单元，而用户、权限和审批可以在边界稳定前继续保留在单体中。'],
'第12章 结论与展望': ['系统设计的核心价值在于把算法能力转化为可管理的质量流程：每张图片有来源，每次检测有运行版本，每个结果有审批人，每次冲突都有处理记录。', '后续工作将围绕数据质量、模型评估和协作效率展开，并在真实使用数据基础上调整工作单元大小、重复判定策略和 Agent 工具白名单。']}
detail_map['第1章 绪论'] += ['系统的研究价值体现在把算法、文件、人员和审批统一为可追溯链路。对于工业企业而言，检测结果不仅是分类标签，还必须说明由谁在何时使用哪个模型对哪批图片产生。', '本研究限定为同一公司检测部门，避免引入与当前规模不匹配的复杂平台能力，同时通过租户字段、事件契约和模块 API 为未来扩展留下接口。']
detail_map['第2章 需求分析与业务规则'] += ['本地电脑上传不能假设员工之间实时在线。系统允许重复上传，再通过指纹和重复工单把不确定性转化为可处理的人工流程。', '5000 张是管理和恢复边界。工作单元过大时失败重试和人员交接成本高，过小时增加派发和审计开销，因此采用最多 5000 张。']
detail_map['第3章 系统总体架构设计'] += ['模块依赖方向是架构质量的关键。检测模块调用基础设施和权限 API，AI 模块不依赖检测数据库，从而可以替换模型或 Agent 运行时。', 'REST 适合短请求，Kafka 适合长耗时推理，AG-UI 适合流式对话。三者分工使上传、推理、审批和对话拥有合理的超时、重试和监控策略。']
detail_map['第4章 检测业务流程与权限模型'] += ['公共可见性由主管审批控制，是最重要的业务闸门。审批前模型结果只能作为员工工作材料，审批后的结果才进入部门知识和统计。', '权限判断必须在 Controller、Service 和 Tool 三层一致。页面隐藏按钮只能改善体验，服务端资源查询和状态修改才是安全边界。']
detail_map['第5章 多用户共享数据与一致性设计'] += ['版本冲突是系统主动保护数据的信号。用户重新加载并合并可以保留双方业务意图，自动覆盖则无法恢复被覆盖意见。', 'Outbox 解决数据库提交和事件发布之间的间隙，发布失败时业务数据仍可查询，发布器可以重试。']
detail_map['第6章 AgentScope 多智能体设计'] += ['专用 Agent 边界清晰后更容易评估：分析 Agent 关注数字准确性，知识 Agent 关注引用命中率，报告 Agent 关注事实完整性。', '安全路由 Agent 不能替代业务权限，而是第二道防线；工具层仍必须拒绝越权。']
detail_map['第7章 数据库与接口设计'] += ['图片指纹表保存最近状态，不等于保存全部历史。历史通过 detection_run 和结果版本关联，指纹只负责快速判断。', '幂等设计贯穿上传确认、工作单元提交和重复工单处理，客户端超时重试不会创建重复记录。']
detail_map['第8章 前端交互与复查对比设计'] += ['双面板把两个来源置于同一视觉坐标，降低主管在标签页之间切换导致的遗漏。同步索引只同步图片位置，业务结论仍需明确点击。', '组件需处理网络慢、签名 URL 过期和数量不一致，图片加载失败不能被误认为检测结果为空。']
detail_map['第9章 部署配置与可观测性'] += ['生产环境密钥由 Secret 管理，配置文件只描述变量映射。没有外部服务时，项目仍可启动基础页面和接口。', '统一 traceId 将上传、推理、审批和公共发布串联，便于定位慢请求、错误结果和权限拒绝。']
detail_map['第10章 测试方案与质量保证'] += ['测试数据应覆盖正常、失败、重复、冲突和越权五类路径。重复工单还要验证来源任务保留和主记录选择可追溯。', '性能优化优先采用索引、分页和异步处理，不能通过降低审计或权限校验换取速度。']
detail_map['第11章 系统演进与微服务拆分'] += ['微服务拆分的收益来自独立扩展和治理，而不是服务数量。当前保持单体可共享事务和调试链路，推理或 Agent 成为瓶颈后再拆分。', '拆分后需保留相同 DTO、错误码和事件语义，前端和上层编排无需因部署方式变化重写。']
detail_map['第12章 结论与展望'] += ['系统将一次检测定义为由图片、运行、操作者、模型、结果和审批组成的完整记录，为质量追溯和责任认定提供数据基础。', '主管对比结论、员工修订意见和工具拒绝记录可以形成评估样本，但必须脱敏并经过权限控制。']
detail_map['第1章 绪论'] += ['与只关注准确率的算法论文不同，本文进一步讨论工程系统中的数据责任、权限边界和异常恢复。研究对象是真实生产中持续产生的图片批次，因此需要同时考虑网络、存储和人员协作的不确定性。']
detail_map['第2章 需求分析与业务规则'] += ['需求被转换为可验证的状态和权限规则，而不是停留在页面描述层面。例如，公共记录的可见条件是审批状态为 APPROVED 或 PUBLIC，员工写入条件是任务所有者和任务状态同时满足。']
detail_map['第3章 系统总体架构设计'] += ['采用模块化单体还可以降低调试成本。一次业务请求在同一进程内完成身份校验、事务更新和审计记录，开发阶段可以快速定位问题；稳定后再通过契约将高负载模块外置。']
detail_map['第4章 检测业务流程与权限模型'] += ['在员工提交审批前，系统应检查图片数量、模型版本、失败数量和必要说明是否完整。主管审批页面展示来源员工、检测时间和结果版本，避免只凭一个汇总数字做决定。']
detail_map['第5章 多用户共享数据与一致性设计'] += ['对于重复图片，系统不试图猜测哪位员工正确，而是保留两份结果并引入对比工单。该策略把技术冲突转化为业务可解释的复核流程，也为责任审计提供证据。']
detail_map['第6章 AgentScope 多智能体设计'] += ['Agent 的系统提示词只描述角色和表达规则，不能承担最终授权。工具返回的事实应带来源和版本，编排器据此生成回答；当事实不足时必须返回待补充信息而不是推测。']
detail_map['第7章 数据库与接口设计'] += ['表结构中的创建人、审批人和姓名快照解决人员变更后的历史追溯问题，version 和状态字段解决并发问题，sha256 和 objectKey 解决图片识别和存储定位问题。']
detail_map['第8章 前端交互与复查对比设计'] += ['主管对比操作需要把图片加载、结果信息和结论表单放在同一个工作上下文中。这样可以减少在多个页面复制文件名和结果的操作，降低人工误判概率。']
detail_map['第9章 部署配置与可观测性'] += ['部署文档应同时说明启动顺序、健康检查、迁移策略、日志位置和回滚方法。对于 AgentScope，缺少 Key 时关闭 Agent Bean，避免开发环境因为外部服务不可用而无法启动。']
detail_map['第10章 测试方案与质量保证'] += ['验收不能只验证成功路径。应模拟重复点击、浏览器断线、Kafka 延迟、OSS 签名过期、数据库版本冲突和权限变化，验证系统能提示用户并保持数据可恢复。']
detail_map['第11章 系统演进与微服务拆分'] += ['事件版本采用 v1 后，新增字段应保持消费者兼容；破坏性变更使用新事件版本并保留一段迁移期。该策略使单体和微服务部署都能逐步演进。']
detail_map['第12章 结论与展望'] += ['最终评价指标应同时包括检测效率、审批时延、重复处理率、人工修订率和 Agent 工具拒绝率。只有算法、流程和协作指标共同改善，系统才真正产生企业价值。']

for title, sections in chapters:
    heading(title)
    for st, text in sections:
        heading(st,2); para(text)
    for extra in detail_map.get(title, []): para(extra)
    if title != '第12章 结论与展望': page()

heading('附录A 关键状态与事件字典')
para('任务状态、工作单元状态和重复工单状态共同构成检测业务的可追踪状态机。状态变更由领域 Service 产生事件，事件包含 eventId、aggregateId、operatorId、occurredAt 和 schemaVersion。')
tbl(['事件','触发条件','下游用途'],[['detection.task.submitted','员工提交检测结果','主管审核池刷新'],['detection.review.approved','主管审批通过','公共记录发布'],['detection.image.duplicate','指纹发现重复','创建重复工单'],['detection.work-unit.recalled','工作单元被回收','重新派发']])
page()
heading('附录B 员工与主管接口示例')
para('员工预检查接口接收图片文件名、对象 Key、SHA-256 和检测策略，返回每张图片的历史状态。主管重复工单详情接口返回左右来源批次、员工信息、图片列表和版本号。写接口统一使用 requestId、expectedVersion 和 confirmationToken，服务端从登录上下文获取操作者身份。')
para('在实际部署中，前端不应缓存长期图片地址。对象存储签名 URL 使用短时有效期，页面刷新或切换图片时重新获取授权地址。')
page()
heading('附录C 安全与运维检查清单')
bullets(['生产环境密钥由 Secret 管理，不提交 .env。','数据库迁移执行前备份并记录版本。','工作区目录使用专用路径和容量配额。','Kafka、Redis、OSS 故障时提供可重试和降级提示。','每月抽查 Agent 审计、工具拒绝和版本冲突记录。'])
page()
heading('附录D 未来实验方向')
para('后续可评估视觉模型置信度校准、重复图片感知哈希阈值、基于历史复核的检索增强、Agent 工具调用准确率和主管审批耗时。实验结果应使用脱敏数据集和固定评估指标，不能直接将实验模型发布到生产检测链路。')
page()
heading('参考文献')
for r in ['[1] AgentScope Java v2 Documentation. https://java.agentscope.io/v2/llms.txt','[2] AgentScope Java Quickstart. https://java.agentscope.io/v2/en/docs/quickstart.html','[3] AgentScope Spark Design. https://github.com/agentscope-ai/agentscope-spark-design','[4] Spring Boot Reference Documentation. https://docs.spring.io/spring-boot/docs/current/reference/html/','[5] Martin Fowler. Patterns of Enterprise Application Architecture. Addison-Wesley, 2002.','[6] Eric Evans. Domain-Driven Design. Addison-Wesley, 2003.','[7] Chris Richardson. Microservices Patterns. Manning, 2018.']: para(r)

for section in doc.sections:
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run('集装箱门把手智能检测系统  |  ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
out.parent.mkdir(parents=True, exist_ok=True); doc.save(out); print(out)
